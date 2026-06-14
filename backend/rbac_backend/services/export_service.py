import asyncio
import io
import logging
from collections.abc import Mapping, Sequence
from datetime import date, datetime
from typing import Any, Dict, List, Optional, Tuple

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExportService:
    """
    Helper service responsible for turning document payloads into binary Excel files.
    The API remains async to align with router usage, but the heavy lifting is
    performed in a background thread to avoid blocking the event loop.
    """

    _COLUMN_DEFINITIONS: Tuple[Tuple[str, Tuple[str, ...], bool], ...] = (
        ("Date", ("date_display", "date"), False),
        ("Letter No.", ("letterNo", "letter_no"), False),
        ("Direction", ("direction", "uploadType"), False),
        ("From", ("from_display", "from", "from_"), False),
        ("To", ("to_display", "to"), False),
        ("Subject", ("subject", "Subject"), False),
        ("Tag", ("tag_display", "tags"), True),
        ("Sub-Tag", ("subTag_display", "subTags"), True),
        ("Status", ("status",), False),
        (
            "Organisation",
            ("organization_name", "organization", "organizationName", "organization_id", "organizationId"),
            False,
        ),
        ("Project", ("project_name", "project", "projectName", "project_id", "projectId"), False),
        ("Upload Date", ("upload_date_display", "uploadDate", "createdAt"), False),
        ("Creater/Uploader Name", ("uploader_name", "createdByName", "createdBy"), False),
    )

    async def export_documents_to_xlsx(
        self,
        documents: List[Dict[str, Any]],
        filename: Optional[str] = None,
    ) -> bytes:
        """
        Convert a list of document dictionaries into an XLSX file represented as bytes.
        """
        docs = documents or []
        try:
            return await asyncio.to_thread(self._build_workbook_bytes, docs, filename)
        except Exception:
            logger.exception("Failed to export documents to XLSX")
            return b""

    def _build_workbook_bytes(
        self,
        documents: List[Dict[str, Any]],
        filename: Optional[str],
    ) -> bytes:
        workbook = Workbook()
        if filename:
            workbook.properties.title = filename

        worksheet = workbook.active
        worksheet.title = "Documents"

        headers = [column[0] for column in self._COLUMN_DEFINITIONS]
        worksheet.append(headers)

        for raw_doc in documents:
            doc = self._ensure_mapping(raw_doc)
            row = [
                self._prepare_cell_value(self._resolve_value(doc, keys), treat_as_sequence)
                for _, keys, treat_as_sequence in self._COLUMN_DEFINITIONS
            ]
            worksheet.append(row)

        self._auto_size_columns(worksheet)

        with io.BytesIO() as buffer:
            workbook.save(buffer)
            buffer.seek(0)
            return buffer.getvalue()

    @staticmethod
    def _ensure_mapping(document: Any) -> Dict[str, Any]:
        if document is None:
            return {}
        if isinstance(document, Mapping):
            return dict(document)
        try:
            if hasattr(document, "model_dump"):
                return document.model_dump(by_alias=True)
        except Exception:
            logger.debug("Failed to model_dump document during export", exc_info=True)
        return {}

    @staticmethod
    def _resolve_value(document: Dict[str, Any], keys: Tuple[str, ...]) -> Any:
        for key in keys:
            if key in document and document[key] not in (None, ""):
                return document[key]
        return None

    @staticmethod
    def _prepare_cell_value(value: Any, treat_as_sequence: bool) -> Any:
        if treat_as_sequence:
            if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
                cleaned = [str(item) for item in value if item not in (None, "")]
                return "; ".join(cleaned)
            if value in (None, ""):
                return ""
            return str(value)

        if isinstance(value, datetime):
            return value.strftime("%d-%m-%Y %H:%M")

        if isinstance(value, date):
            return value.strftime("%d-%m-%Y")

        if isinstance(value, Mapping):
            return str(dict(value))

        if isinstance(value, Sequence) and not isinstance(value, (str, bytes)):
            return "; ".join(str(item) for item in value if item not in (None, ""))

        if value in (None, ""):
            return ""

        return value

    @staticmethod
    def _auto_size_columns(worksheet) -> None:
        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)
            for cell in column_cells:
                if cell.value is None:
                    continue
                cell_length = len(str(cell.value))
                if cell_length > max_length:
                    max_length = cell_length
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 60)
