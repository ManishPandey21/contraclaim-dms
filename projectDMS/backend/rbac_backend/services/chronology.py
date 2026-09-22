from __future__ import annotations

import csv
import hashlib
import io
import re
from datetime import datetime
from enum import Enum
from typing import Any, Dict, List, Optional
from xml.sax.saxutils import escape

from fastapi import HTTPException, status

from ..models.arbitration_drafting import (
    ArbitrationSelectedReference,
    ArbitrationSourceType,
    ArbitrationSourceUse,
)
from ..models.chronology import (
    AttachChronologyRequest,
    ChronologyDateType,
    ChronologyDecisionRequest,
    ChronologyDuplicateRequest,
    ChronologyEventClassification,
    ChronologyExtractRequest,
    ChronologyImpactType,
    ChronologyLinkRequest,
    ChronologyPleadingContext,
    ChronologyRevisionAction,
    ChronologyStatus,
    ChronologyVerificationStatus,
    MatterChronology,
    MatterChronologyCreate,
    MatterChronologyEvent,
    MatterChronologyEventCreate,
    MatterChronologyEventRevision,
    MatterChronologyEventUpdate,
    MatterChronologyUpdate,
)
from ..models.evidence_graph import (
    AIExtractionCreate,
    EventLinkCreate,
    EventLinkStatus,
    EventRelationType,
    EvidenceEntityType,
    ProjectEventCreate,
    ProjectEventStatus,
    ProjectEventType,
    SourceSpan,
)
from .arbitration_drafting.repository import ArbitrationDraftingRepository
from .audit_event_service import AuditEventService
from .evidence_graph_service import EvidenceGraphService


def _actor_id(user: Any) -> Optional[str]:
    return getattr(user, "id", None) or getattr(user, "email", None)


#: The chronology-event fields that carry AUTHORITY rather than metadata.
#:
#: `PolicyService.authorize_document` resolves a resource's authority from
#: `organization_id` and `project_id` and nothing else, so these two — and only
#: these two — decide which tenant an event, its `project_events` row, its
#: `event_links` and its `audit_events` land in, and which tenant's
#: `build_scope_query` reads them back. `contract_id` / `matter_id` / `claim_id`
#: are list-filter metadata that no gate reads; `MatterChronologyUpdate` already
#: lets a chronology change all three under an org/project-only gate, so they
#: stay caller-settable.
_EVENT_AUTHORITY_FIELDS = ("organization_id", "project_id")


def _enum_value(value: Any) -> Any:
    return value.value if isinstance(value, Enum) else value


def _jsonable(value: Any) -> Any:
    if isinstance(value, Enum):
        return value.value
    if isinstance(value, datetime):
        return value
    if isinstance(value, dict):
        return {key: _jsonable(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_jsonable(item) for item in value]
    return value


async def _collect(cursor: Any) -> List[Dict[str, Any]]:
    if hasattr(cursor, "to_list"):
        return [dict(item) for item in await cursor.to_list(length=None)]
    return [dict(item) async for item in cursor]


def _shorten(text: Any, width: int = 700) -> str:
    raw = " ".join(str(text or "").split())
    if len(raw) <= width:
        return raw
    return raw[: width - 3].rstrip() + "..."


def _source_hash(row: Dict[str, Any]) -> str:
    raw = "|".join(
        [
            str(row.get("source_type") or ""),
            str(row.get("source_id") or ""),
            str(row.get("citation") or ""),
            str(row.get("snippet") or ""),
        ]
    )
    return hashlib.sha256(raw.encode("utf-8")).hexdigest()


class ChronologyService:
    def __init__(self, db: Any) -> None:
        self.db = db
        self.audit = AuditEventService(db)
        self.graph = EvidenceGraphService(db)
        self.arbitration_repo = ArbitrationDraftingRepository(db)

    async def create_chronology(self, payload: MatterChronologyCreate, current_user: Any) -> Dict[str, Any]:
        data = payload.model_dump()
        data["organization_id"] = data.get("organization_id") or getattr(current_user, "organization_id", None)
        doc = MatterChronology(
            **data,
            created_by=_actor_id(current_user),
            updated_at=datetime.utcnow(),
        ).model_dump(by_alias=True, exclude_none=True)
        await self.db.matter_chronologies.insert_one(_jsonable(doc))
        await self._emit("created", doc, current_user, after=doc)
        return await self.get_chronology(doc["_id"])

    async def list_chronologies(
        self,
        scope: Dict[str, Any],
        filters: Dict[str, Any],
        *,
        skip: int = 0,
        limit: int = 100,
    ) -> List[Dict[str, Any]]:
        query = dict(scope or {})
        query["deleted_at"] = {"$exists": False}
        for key in ["project_id", "contract_id", "matter_id", "claim_id", "chronology_type", "party_perspective", "status"]:
            if filters.get(key):
                query[key] = filters[key]
        if filters.get("q"):
            query["title"] = {"$regex": str(filters["q"]), "$options": "i"}
        cursor = self.db.matter_chronologies.find(query).sort("updated_at", -1).skip(skip).limit(limit)
        return await _collect(cursor)

    async def get_chronology(self, chronology_id: str) -> Dict[str, Any]:
        doc = await self.db.matter_chronologies.find_one({"_id": chronology_id, "deleted_at": {"$exists": False}})
        if not doc:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chronology not found")
        return dict(doc)

    async def update_chronology(self, chronology_id: str, payload: MatterChronologyUpdate, current_user: Any) -> Dict[str, Any]:
        before = await self.get_chronology(chronology_id)
        update = payload.model_dump(exclude_unset=True)
        update["updated_at"] = datetime.utcnow()
        update["updated_by"] = _actor_id(current_user)
        updated = await self.db.matter_chronologies.find_one_and_update(
            {"_id": chronology_id},
            {"$set": _jsonable(update)},
            return_document=True,
        )
        await self._emit("updated", before, current_user, before=before, after=updated)
        return dict(updated)

    async def delete_chronology(
        self, chronology_id: str, current_user: Any, *, relationships: Any = None
    ) -> None:
        """Soft-delete a chronology.

        With ``relationships`` (a ``DocumentRelationshipService``) the canonical
        Document links of its events are retired in the same transaction, so no
        active link is left pointing at an event of a deleted chronology (CL-3B).
        The events and every Document stay untouched.
        """
        doc = await self.get_chronology(chronology_id)

        async def soft_delete(session: Any = None) -> None:
            now = datetime.utcnow()
            kwargs = {"session": session} if session is not None else {}
            result = await self.db.matter_chronologies.update_one(
                {"_id": chronology_id, "deleted_at": {"$exists": False}},
                {"$set": {"deleted_at": now, "deleted_by": _actor_id(current_user), "updated_at": now}},
                **kwargs,
            )
            if not getattr(result, "matched_count", 0):
                raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail="Chronology changed during deletion")

        if relationships is None:
            await soft_delete()
        else:

            async def event_ids(session: Any) -> List[str]:
                kwargs = {"session": session} if session is not None else {}
                cursor = self.db.matter_chronology_events.find({"chronology_id": chronology_id}, {"_id": 1}, **kwargs)
                return [str(row.get("_id")) async for row in cursor if row.get("_id")]

            await relationships.retire_target_links(
                current_user,
                "chronology_event",
                organization_id=str(doc.get("organization_id") or ""),
                project_id=str(doc.get("project_id") or ""),
                target_ids=event_ids,
                reason="Chronology deleted",
                before=soft_delete,
            )
        await self._emit("deleted", doc, current_user)

    async def list_events(
        self,
        chronology_id: str,
        filters: Dict[str, Any],
        *,
        skip: int = 0,
        limit: int = 250,
    ) -> List[Dict[str, Any]]:
        await self.get_chronology(chronology_id)
        query: Dict[str, Any] = {"chronology_id": chronology_id}
        for key in [
            "verification_status",
            "supports_party",
            "event_classification",
            "impact_type",
            "responsible_party",
            "pleading_use",
        ]:
            if filters.get(key):
                query[key] = filters[key]
        if filters.get("issue_tag"):
            query["issue_tags"] = filters["issue_tag"]
        if filters.get("claim_head"):
            query["claim_heads"] = filters["claim_head"]
        if filters.get("clause"):
            query["contract_clauses"] = {"$regex": re.escape(filters["clause"]), "$options": "i"}
        if filters.get("source_type"):
            query["metadata.source_type"] = filters["source_type"]
        if filters.get("date_from") or filters.get("date_to"):
            range_query: Dict[str, Any] = {}
            if filters.get("date_from"):
                range_query["$gte"] = filters["date_from"]
            if filters.get("date_to"):
                range_query["$lte"] = filters["date_to"]
            query["event_date"] = range_query
        cursor = self.db.matter_chronology_events.find(query).sort([("event_date", 1), ("created_at", 1)]).skip(skip).limit(limit)
        events = await _collect(cursor)
        # `description`/`title`/`source_spans` are span text lifted from the
        # source document (`chronology.py:530`). One shared projection withholds
        # them when the originating document is no longer consumable; manual
        # events (no source_document_id) are untouched.
        from .publication_policy import safe_event_records

        return await safe_event_records(
            self.db, events, ("description", "title"), span_fields=("source_spans",)
        )

    def _apply_chronology_authority_scope(self, data: Dict[str, Any], chronology: Dict[str, Any]) -> None:
        """Anchor an event's authority scope to its already-authorized parent.

        The route authorizes the PARENT CHRONOLOGY — `_load_and_authorize_chronology`
        hands the policy the chronology row, so nothing in the request body is ever
        authorized. Resolving scope as `payload_value or chronology_value` therefore
        let a caller redirect the event, and everything derived from it, into a
        tenant the request was never checked against.

        Authority is the parent chronology intersected with the actor's entitlement,
        and the route has already established the second half. A request field may
        restate that scope but never replace it: a conflicting value is refused
        rather than silently rewritten, so a caller that believes it is writing
        somewhere else finds out. Actor entitlement is not parent identity — a
        globally entitled actor is refused here too, because this API has no
        relocation semantics.
        """
        for field in _EVENT_AUTHORITY_FIELDS:
            authorized = chronology.get(field)
            requested = data.get(field)
            if requested is not None and str(requested) != str(authorized):
                raise HTTPException(
                    status_code=status.HTTP_403_FORBIDDEN,
                    detail=(
                        f"Chronology event {field} must match its chronology; "
                        "an event cannot be created outside the authorized chronology scope."
                    ),
                )
            data[field] = authorized

    async def create_event(self, payload: MatterChronologyEventCreate, current_user: Any) -> Dict[str, Any]:
        chronology = await self.get_chronology(payload.chronology_id)
        if payload.source_document_id:
            await self._require_current_document_authority(payload.source_document_id)
        data = payload.model_dump()
        self._apply_chronology_authority_scope(data, chronology)
        data["contract_id"] = data.get("contract_id") or chronology.get("contract_id")
        data["matter_id"] = data.get("matter_id") or chronology.get("matter_id")
        data["claim_id"] = data.get("claim_id") or chronology.get("claim_id")
        doc = MatterChronologyEvent(
            **data,
            created_by=_actor_id(current_user),
            updated_at=datetime.utcnow(),
        ).model_dump(by_alias=True, exclude_none=True)
        await self.db.matter_chronology_events.insert_one(_jsonable(doc))
        await self._append_revision(doc, ChronologyRevisionAction.CREATED, current_user, after=doc)
        await self._refresh_counts(payload.chronology_id)
        return await self.get_event(payload.chronology_id, doc["_id"])

    async def get_event(self, chronology_id: str, event_id: str) -> Dict[str, Any]:
        doc = await self.db.matter_chronology_events.find_one({"_id": event_id, "chronology_id": chronology_id})
        if not doc:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Chronology event not found")
        return dict(doc)

    async def update_event(
        self,
        chronology_id: str,
        event_id: str,
        payload: MatterChronologyEventUpdate,
        current_user: Any,
    ) -> Dict[str, Any]:
        before = await self.get_event(chronology_id, event_id)
        update = payload.model_dump(exclude_unset=True)
        if before.get("verification_status") in {
            ChronologyVerificationStatus.VERIFIED.value,
            ChronologyVerificationStatus.EDITED_VERIFIED.value,
        } and "verification_status" not in update:
            update["verification_status"] = ChronologyVerificationStatus.EDITED_VERIFIED.value
        update["updated_at"] = datetime.utcnow()
        update["updated_by"] = _actor_id(current_user)
        updated = await self.db.matter_chronology_events.find_one_and_update(
            {"_id": event_id, "chronology_id": chronology_id},
            {"$set": _jsonable(update)},
            return_document=True,
        )
        await self._append_revision(updated, ChronologyRevisionAction.EDITED, current_user, before=before, after=updated)
        await self._refresh_counts(chronology_id)
        return dict(updated)

    async def verify_event(
        self,
        chronology_id: str,
        event_id: str,
        current_user: Any,
        decision: Optional[ChronologyDecisionRequest] = None,
    ) -> Dict[str, Any]:
        before = await self.get_event(chronology_id, event_id)
        if before.get("source_document_id"):
            await self._require_current_document_authority(before["source_document_id"])
        update = {
            "verification_status": ChronologyVerificationStatus.VERIFIED.value,
            "updated_at": datetime.utcnow(),
            "updated_by": _actor_id(current_user),
        }
        updated = await self.db.matter_chronology_events.find_one_and_update(
            {"_id": event_id, "chronology_id": chronology_id},
            {"$set": update},
            return_document=True,
        )
        synced = await self._sync_verified_event(dict(updated), current_user)
        await self._append_revision(
            synced,
            ChronologyRevisionAction.VERIFIED,
            current_user,
            before=before,
            after=synced,
            note=decision.note if decision else None,
        )
        await self._refresh_counts(chronology_id)
        return synced

    async def reject_event(
        self,
        chronology_id: str,
        event_id: str,
        current_user: Any,
        decision: Optional[ChronologyDecisionRequest] = None,
    ) -> Dict[str, Any]:
        before = await self.get_event(chronology_id, event_id)
        updated = await self.db.matter_chronology_events.find_one_and_update(
            {"_id": event_id, "chronology_id": chronology_id},
            {
                "$set": {
                    "verification_status": ChronologyVerificationStatus.REJECTED.value,
                    "updated_at": datetime.utcnow(),
                    "updated_by": _actor_id(current_user),
                }
            },
            return_document=True,
        )
        await self._append_revision(
            updated,
            ChronologyRevisionAction.REJECTED,
            current_user,
            before=before,
            after=updated,
            note=decision.note if decision else None,
        )
        await self._refresh_counts(chronology_id)
        return dict(updated)

    async def mark_duplicate(
        self,
        chronology_id: str,
        event_id: str,
        payload: ChronologyDuplicateRequest,
        current_user: Any,
    ) -> Dict[str, Any]:
        before = await self.get_event(chronology_id, event_id)
        await self.get_event(chronology_id, payload.duplicate_of_event_id)
        updated = await self.db.matter_chronology_events.find_one_and_update(
            {"_id": event_id, "chronology_id": chronology_id},
            {
                "$set": {
                    "verification_status": ChronologyVerificationStatus.DUPLICATE.value,
                    "duplicate_of_event_id": payload.duplicate_of_event_id,
                    "updated_at": datetime.utcnow(),
                    "updated_by": _actor_id(current_user),
                }
            },
            return_document=True,
        )
        await self._append_revision(
            updated,
            ChronologyRevisionAction.MARKED_DUPLICATE,
            current_user,
            before=before,
            after=updated,
            note=payload.note,
        )
        await self._refresh_counts(chronology_id)
        return dict(updated)

    async def link_event(
        self,
        chronology_id: str,
        event_id: str,
        payload: ChronologyLinkRequest,
        current_user: Any,
    ) -> Dict[str, Any]:
        before = await self.get_event(chronology_id, event_id)
        related_event_ids = sorted(set((before.get("related_event_ids") or []) + payload.related_event_ids))
        related_document_ids = sorted(set((before.get("related_document_ids") or []) + payload.related_document_ids))
        updated = await self.db.matter_chronology_events.find_one_and_update(
            {"_id": event_id, "chronology_id": chronology_id},
            {
                "$set": {
                    "related_event_ids": related_event_ids,
                    "related_document_ids": related_document_ids,
                    "updated_at": datetime.utcnow(),
                    "updated_by": _actor_id(current_user),
                }
            },
            return_document=True,
        )
        await self._append_revision(updated, ChronologyRevisionAction.LINKED, current_user, before=before, after=updated, note=payload.note)
        return dict(updated)

    async def event_revisions(self, chronology_id: str, event_id: str) -> List[Dict[str, Any]]:
        await self.get_event(chronology_id, event_id)
        cursor = self.db.matter_chronology_event_revisions.find({"chronology_id": chronology_id, "event_id": event_id}).sort("revision", 1)
        revisions = await _collect(cursor)
        # `before`/`after` are event snapshots embedding document-derived span
        # text. Gate the nested snapshots through the same projection so the
        # audit surface cannot serve a blocked document's text either.
        from .publication_policy import safe_event_records

        for rev in revisions:
            for key in ("before", "after"):
                snap = rev.get(key)
                if isinstance(snap, dict):
                    rev[key] = (
                        await safe_event_records(
                            self.db, [snap], ("description", "title"), span_fields=("source_spans",)
                        )
                    )[0]
        return revisions

    async def extract_events(self, chronology_id: str, payload: ChronologyExtractRequest, current_user: Any) -> Dict[str, Any]:
        chronology = await self.get_chronology(chronology_id)
        await self.db.matter_chronologies.update_one(
            {"_id": chronology_id},
            {"$set": {"status": ChronologyStatus.EXTRACTING.value, "updated_at": datetime.utcnow(), "updated_by": _actor_id(current_user)}},
        )
        source_ids = payload.source_document_ids or chronology.get("selected_source_ids") or []
        query: Dict[str, Any] = {
            "organization_id": chronology.get("organization_id"),
            "project_id": chronology.get("project_id"),
        }
        if source_ids:
            from .publication_policy import document_id_candidates

            _ids = []
            for _sid in source_ids[: payload.max_documents]:
                _ids.extend(document_id_candidates(_sid))
            query["_id"] = {"$in": _ids}
        cursor = self.db.documents.find(query).limit(payload.max_documents)
        docs = await _collect(cursor)
        created: List[Dict[str, Any]] = []
        for source_doc in docs:
            event = await self._extract_document_event(chronology, source_doc, current_user)
            if event:
                created.append(event)
        await self.db.matter_chronologies.update_one(
            {"_id": chronology_id},
            {"$set": {"status": ChronologyStatus.REVIEW.value, "updated_at": datetime.utcnow()}},
        )
        await self._refresh_counts(chronology_id)
        await self._emit("extracted", chronology, current_user, after={"created": len(created), "documents": len(docs)})
        return {"chronology_id": chronology_id, "documents_scanned": len(docs), "events_created": len(created), "events": created}

    async def pleading_context(
        self,
        chronology_id: str,
        *,
        include_unverified: bool = False,
        limit: int = 250,
    ) -> ChronologyPleadingContext:
        chronology = await self.get_chronology(chronology_id)
        statuses = [ChronologyVerificationStatus.VERIFIED.value, ChronologyVerificationStatus.EDITED_VERIFIED.value]
        if include_unverified:
            statuses.extend([ChronologyVerificationStatus.AI_SUGGESTED.value, ChronologyVerificationStatus.NEEDS_REVIEW.value])
        cursor = (
            self.db.matter_chronology_events.find({"chronology_id": chronology_id, "verification_status": {"$in": statuses}})
            .sort([("event_date", 1), ("created_at", 1)])
            .limit(limit)
        )
        events = await _collect(cursor)
        # Same projection as `list_events`: a document-derived event's
        # description/title/source_spans are span text lifted from the source
        # document, so they must be withheld once that document is no longer
        # consumable. `_ledger_row` reads exactly those fields, so gate before
        # building the ledger. Manual events (no source_document_id) untouched.
        from .publication_policy import safe_event_records

        events = await safe_event_records(
            self.db, events, ("description", "title"), span_fields=("source_spans",)
        )
        ledger = [self._ledger_row(event, idx) for idx, event in enumerate(events, start=1)]
        missing = []
        for event in events:
            if not event.get("source_document_id") and not event.get("source_spans"):
                missing.append(f"Chronology event needs source backing: {event.get('title')}")
            if event.get("verification_status") not in {ChronologyVerificationStatus.VERIFIED.value, ChronologyVerificationStatus.EDITED_VERIFIED.value}:
                missing.append(f"Unverified chronology event included for review only: {event.get('title')}")
        return ChronologyPleadingContext(
            chronology_id=chronology_id,
            source_ledger=ledger,
            missing_evidence=missing,
            summary={"events": len(events), **(chronology.get("summary_counts") or {})},
        )

    async def attach_to_arbitration_draft(
        self,
        draft_id: str,
        payload: AttachChronologyRequest,
        current_user: Any,
    ) -> Dict[str, Any]:
        draft = await self.arbitration_repo.get_draft(draft_id)
        if not draft:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Arbitration draft not found")
        chronology = await self.get_chronology(payload.chronology_id)
        if draft.get("organization_id") != chronology.get("organization_id") or draft.get("project_id") != chronology.get("project_id"):
            raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Chronology is outside the arbitration draft scope")
        context = await self.pleading_context(payload.chronology_id, include_unverified=payload.include_unverified, limit=payload.limit)
        refs = [
            ArbitrationSelectedReference(
                draft_id=draft_id,
                source_type=ArbitrationSourceType.CHRONOLOGY_EVENT,
                source_id=row["source_id"],
                label=row["label"],
                citation=row.get("citation"),
                snippet=row.get("snippet"),
                page_numbers=row.get("page_numbers") or [],
                clause_number=row.get("clause_number"),
                letter_no=row.get("letter_no"),
                allowed_use=ArbitrationSourceUse.CHRONOLOGY,
                selected_by=_actor_id(current_user),
                metadata=row.get("metadata") or {},
            ).model_dump(by_alias=True)
            for row in context.source_ledger
        ]
        saved = await self.arbitration_repo.add_references(draft_id, refs)
        await self._emit("attached_to_arbitration", chronology, current_user, after={"draft_id": draft_id, "reference_count": len(saved)})
        return {"draft_id": draft_id, "chronology_id": payload.chronology_id, "reference_count": len(saved), "missing_evidence": context.missing_evidence}

    async def export(self, chronology_id: str, fmt: str, current_user: Any) -> bytes:
        chronology = await self.get_chronology(chronology_id)
        events = await self.list_events(chronology_id, {}, limit=1000)
        if fmt == "xlsx":
            content = self._csv_bytes(events)
        elif fmt == "pdf":
            content = self._pdf_bytes(chronology, events)
        else:
            content = self._docx_bytes(chronology, events)
        await self.db.matter_chronologies.update_one(
            {"_id": chronology_id},
            {"$set": {"status": ChronologyStatus.EXPORTED.value, "updated_at": datetime.utcnow(), "updated_by": _actor_id(current_user)}},
        )
        await self._emit("exported", chronology, current_user, after={"format": fmt, "event_count": len(events)})
        return content

    async def _extract_document_event(self, chronology: Dict[str, Any], source_doc: Dict[str, Any], current_user: Any) -> Optional[Dict[str, Any]]:
        source_id = str(source_doc.get("_id"))
        if not await self._has_current_document_authority(source_id):
            return None
        raw_text = self._document_text(source_doc)
        content_hash = hashlib.sha256(raw_text.encode("utf-8")).hexdigest()
        existing = await self.db.matter_chronology_events.find_one(
            {"chronology_id": chronology["_id"], "source_document_id": source_id, "metadata.content_hash": content_hash}
        )
        if existing:
            return dict(existing)
        event_date, date_text, date_type = self._event_date(source_doc, raw_text)
        clauses = self._clauses(raw_text)
        letter_no = source_doc.get("letterNo") or source_doc.get("letter_no") or self._letter_no(raw_text)
        classification = self._classification(raw_text)
        issue_tags = self._issue_tags(raw_text, classification)
        extraction = await self.graph.create_ai_extraction(
            AIExtractionCreate(
                organization_id=chronology.get("organization_id"),
                project_id=chronology.get("project_id"),
                source_document_id=source_id,
                content_hash=content_hash,
                schema_version="chronology.v1",
                model="deterministic-rule-extractor",
                prompt_version="chronology.v1",
                raw_output={"text_sample": _shorten(raw_text, 1200)},
                parsed_output={"event_date": date_text, "clauses": clauses, "letter_no": letter_no, "classification": classification.value},
                confidence=0.72 if raw_text and event_date else 0.45,
            ),
            current_user=current_user,
        )
        span_text = _shorten(raw_text, 900)
        event = MatterChronologyEventCreate(
            chronology_id=chronology["_id"],
            organization_id=chronology.get("organization_id"),
            project_id=chronology.get("project_id"),
            contract_id=chronology.get("contract_id"),
            matter_id=chronology.get("matter_id"),
            claim_id=chronology.get("claim_id"),
            event_date=event_date,
            date_text=date_text,
            date_type=date_type,
            title=self._title(source_doc, raw_text),
            description=span_text or None,
            source_document_id=source_id,
            source_document_name=source_doc.get("filename") or source_doc.get("file_name") or source_doc.get("subject"),
            source_page=self._first_page(source_doc),
            source_spans=[{"page": self._first_page(source_doc), "text": span_text}] if span_text else [],
            letter_no=letter_no,
            from_party=source_doc.get("from_company") or source_doc.get("from") or source_doc.get("sender"),
            to_party=source_doc.get("to_company") or source_doc.get("to") or source_doc.get("recipient"),
            contract_clauses=clauses,
            issue_tags=issue_tags,
            claim_heads=issue_tags,
            event_classification=classification,
            impact_type=self._impact_type(classification),
            confidence_score=0.72 if raw_text and event_date else 0.45,
            verification_status=ChronologyVerificationStatus.AI_SUGGESTED if span_text else ChronologyVerificationStatus.NEEDS_REVIEW,
            pleading_use=self._pleading_use(classification, chronology.get("party_perspective")),
            related_document_ids=[source_id],
            ai_extraction_id=str(extraction.get("_id")),
            metadata={"content_hash": content_hash, "source_type": "document"},
        )
        return await self.create_event(event, current_user)

    async def _has_current_document_authority(self, document_id: str) -> bool:
        """Resolve authority from the canonical document at the write boundary."""
        from .publication_policy import resolve_document_authority

        return (await resolve_document_authority(self.db, document_id)).consumable

    async def _require_current_document_authority(self, document_id: str) -> None:
        if not await self._has_current_document_authority(document_id):
            raise HTTPException(
                status_code=status.HTTP_409_CONFLICT,
                detail="Source document is not currently authoritative",
            )

    def _document_text(self, source_doc: Dict[str, Any]) -> str:
        # Suggested chronology entries feed EOT and claim reasoning, so this is
        # an authoritative-content consumer. Metadata stays available so a
        # blocked document is still identifiable; only extracted body text and
        # its derived summary are withheld.
        from .publication_policy import is_consumable

        consumable = is_consumable(source_doc)
        parts = [
            source_doc.get("subject"),
            source_doc.get("summary") if consumable else None,
            source_doc.get("description"),
            source_doc.get("ocrText") if consumable else None,
            source_doc.get("ocr_text") if consumable else None,
            source_doc.get("full_content") if consumable else None,
            source_doc.get("content") if consumable else None,
            source_doc.get("text") if consumable else None,
            source_doc.get("filename"),
        ]
        return " ".join(str(part) for part in parts if part)

    def _event_date(self, source_doc: Dict[str, Any], raw_text: str) -> tuple[Optional[datetime], Optional[str], ChronologyDateType]:
        for key in ["date", "document_date", "letter_date", "created_at"]:
            value = source_doc.get(key)
            parsed = self._parse_date(value)
            if parsed:
                return parsed, parsed.date().isoformat(), ChronologyDateType.EXACT
        match = re.search(r"\b(\d{4}-\d{2}-\d{2}|\d{1,2}[./-]\d{1,2}[./-]\d{2,4})\b", raw_text or "")
        if match:
            parsed = self._parse_date(match.group(1))
            if parsed:
                return parsed, match.group(1), ChronologyDateType.EXACT
        return None, None, ChronologyDateType.UNDATED

    def _parse_date(self, value: Any) -> Optional[datetime]:
        if isinstance(value, datetime):
            return value
        if not value:
            return None
        raw = str(value).strip()
        for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%y", "%d/%m/%y", "%d-%m-%y"):
            try:
                return datetime.strptime(raw[:10], fmt)
            except ValueError:
                continue
        try:
            return datetime.fromisoformat(raw.replace("Z", "+00:00")).replace(tzinfo=None)
        except ValueError:
            return None

    def _clauses(self, text: str) -> List[str]:
        found = re.findall(r"\b(?:GCC|SCC|Clause|Cl\.?)\s*([0-9]+(?:\.[0-9]+)*)\b", text or "", flags=re.I)
        out = []
        for raw in found:
            label = f"Clause {raw}"
            if label not in out:
                out.append(label)
        return out[:10]

    def _letter_no(self, text: str) -> Optional[str]:
        match = re.search(r"\b[A-Z]{2,5}/\d{4}/\d{3,6}\b", text or "")
        return match.group(0) if match else None

    def _classification(self, text: str) -> ChronologyEventClassification:
        lower = (text or "").lower()
        if any(word in lower for word in ["delay", "eot", "extension of time", "critical path"]):
            return ChronologyEventClassification.DELAY
        if any(word in lower for word in ["payment", "ipc", "invoice", "bill", "certificate"]):
            return ChronologyEventClassification.PAYMENT
        if any(word in lower for word in ["variation", "change order", "change request"]):
            return ChronologyEventClassification.VARIATION
        if any(word in lower for word in ["notice", "notification"]):
            return ChronologyEventClassification.NOTICE
        if any(word in lower for word in ["breach", "default", "non-compliance"]):
            return ChronologyEventClassification.BREACH
        return ChronologyEventClassification.EVIDENCE_ONLY

    def _impact_type(self, classification: ChronologyEventClassification) -> ChronologyImpactType:
        if classification == ChronologyEventClassification.DELAY:
            return ChronologyImpactType.TIME
        if classification in {ChronologyEventClassification.PAYMENT, ChronologyEventClassification.VARIATION}:
            return ChronologyImpactType.COST
        if classification == ChronologyEventClassification.BREACH:
            return ChronologyImpactType.LEGAL
        return ChronologyImpactType.UNKNOWN

    def _issue_tags(self, text: str, classification: ChronologyEventClassification) -> List[str]:
        tags = [classification.value]
        lower = (text or "").lower()
        if "notice" in lower and "notice" not in tags:
            tags.append("notice")
        if "quantum" in lower or "cost" in lower:
            tags.append("quantum")
        return tags

    def _pleading_use(self, classification: ChronologyEventClassification, perspective: Optional[str]) -> str:
        if perspective == "respondent":
            return "sod_defence"
        if classification == ChronologyEventClassification.PAYMENT:
            return "soc_quantum"
        if classification in {ChronologyEventClassification.BREACH, ChronologyEventClassification.DELAY, ChronologyEventClassification.NOTICE}:
            return "soc_breach"
        return "soc_background"

    def _title(self, source_doc: Dict[str, Any], raw_text: str) -> str:
        return _shorten(source_doc.get("subject") or source_doc.get("filename") or raw_text or "Chronology event", 160)

    def _first_page(self, source_doc: Dict[str, Any]) -> Optional[int]:
        pages = source_doc.get("page_numbers") or source_doc.get("pages")
        if isinstance(pages, list) and pages:
            try:
                return int(pages[0])
            except Exception:
                return None
        return None

    async def _event_publication_scope(self, event: Dict[str, Any]) -> tuple[Optional[str], Optional[str]]:
        """The scope a verified event may publish under: its parent chronology's.

        Never the stored row's own fields. `create_project_event` re-authorizes
        nothing, so a row written before `_apply_chronology_authority_scope`
        existed would otherwise still publish a project event, its links and its
        audit trail into whatever tenant it happens to carry.
        """
        chronology = await self.get_chronology(event["chronology_id"])
        return chronology.get("organization_id"), chronology.get("project_id")

    async def _sync_verified_event(self, event: Dict[str, Any], current_user: Any) -> Dict[str, Any]:
        if event.get("project_event_id"):
            return event
        if event.get("source_document_id"):
            await self._require_current_document_authority(event["source_document_id"])
        organization_id, project_id = await self._event_publication_scope(event)
        project_event = await self.graph.create_project_event(
            ProjectEventCreate(
                organization_id=organization_id,
                project_id=project_id,
                event_type=self._project_event_type(event),
                event_date=event.get("event_date") or datetime.utcnow(),
                event_end_date=event.get("event_end_date"),
                title=event.get("title") or "Chronology event",
                description=event.get("description"),
                party=event.get("responsible_party") or event.get("from_party"),
                impact_area=event.get("impact_type"),
                source_entity_type=EvidenceEntityType.DOCUMENT if event.get("source_document_id") else None,
                source_entity_id=event.get("source_document_id"),
                status=ProjectEventStatus.OPEN,
                confidence=event.get("confidence_score"),
                ai_extraction_id=event.get("ai_extraction_id"),
                metadata={
                    "chronology_id": event.get("chronology_id"),
                    "chronology_event_id": event.get("_id"),
                    "date_type": event.get("date_type"),
                    "date_text": event.get("date_text"),
                    "issue_tags": event.get("issue_tags") or [],
                    "claim_heads": event.get("claim_heads") or [],
                    "supports_party": event.get("supports_party"),
                    "pleading_use": event.get("pleading_use"),
                    "clauses": event.get("contract_clauses") or [],
                },
            ),
            current_user,
        )
        link_ids: List[str] = []
        if event.get("source_document_id"):
            await self._require_current_document_authority(event["source_document_id"])
            link = await self.graph.suggest_link(
                EventLinkCreate(
                    organization_id=organization_id,
                    project_id=project_id,
                    source_type=EvidenceEntityType.PROJECT_EVENT,
                    source_id=str(project_event["_id"]),
                    target_type=EvidenceEntityType.DOCUMENT,
                    target_id=str(event.get("source_document_id")),
                    relation_type=EventRelationType.REFERS_TO,
                    status=EventLinkStatus.USER_VERIFIED,
                    confidence=event.get("confidence_score"),
                    evidence_text=event.get("description"),
                    source_spans=[SourceSpan(**span) for span in event.get("source_spans") or []],
                    ai_extraction_id=event.get("ai_extraction_id"),
                    metadata={"chronology_id": event.get("chronology_id"), "chronology_event_id": event.get("_id")},
                ),
                current_user=current_user,
            )
            link_ids.append(str(link.get("link_group_id")))
        for clause in event.get("contract_clauses") or []:
            if event.get("source_document_id"):
                await self._require_current_document_authority(event["source_document_id"])
            link = await self.graph.suggest_link(
                EventLinkCreate(
                    organization_id=organization_id,
                    project_id=project_id,
                    source_type=EvidenceEntityType.PROJECT_EVENT,
                    source_id=str(project_event["_id"]),
                    target_type=EvidenceEntityType.CLAUSE,
                    target_id=clause,
                    relation_type=EventRelationType.GOVERNED_BY,
                    status=EventLinkStatus.USER_VERIFIED,
                    confidence=event.get("confidence_score"),
                    evidence_text=event.get("description"),
                    ai_extraction_id=event.get("ai_extraction_id"),
                    metadata={"chronology_id": event.get("chronology_id"), "chronology_event_id": event.get("_id")},
                ),
                current_user=current_user,
            )
            link_ids.append(str(link.get("link_group_id")))
        if event.get("source_document_id"):
            await self._require_current_document_authority(event["source_document_id"])
        updated = await self.db.matter_chronology_events.find_one_and_update(
            {"_id": event["_id"], "chronology_id": event["chronology_id"]},
            {
                "$set": {
                    "project_event_id": str(project_event["_id"]),
                    "event_link_ids": link_ids,
                    "organization_id": organization_id,
                    "project_id": project_id,
                    "updated_at": datetime.utcnow(),
                }
            },
            return_document=True,
        )
        return dict(updated)

    def _project_event_type(self, event: Dict[str, Any]) -> ProjectEventType:
        mapping = {
            ChronologyEventClassification.DELAY.value: ProjectEventType.DELAY,
            ChronologyEventClassification.PAYMENT.value: ProjectEventType.PAYMENT,
            ChronologyEventClassification.VARIATION.value: ProjectEventType.VARIATION,
            ChronologyEventClassification.NOTICE.value: ProjectEventType.LETTER,
            ChronologyEventClassification.BREACH.value: ProjectEventType.CLAIM,
        }
        return mapping.get(event.get("event_classification"), ProjectEventType.OTHER)

    def _ledger_row(self, event: Dict[str, Any], idx: int) -> Dict[str, Any]:
        citation_parts = [event.get("date_text") or (event.get("event_date").date().isoformat() if isinstance(event.get("event_date"), datetime) else None)]
        if event.get("source_document_name"):
            citation_parts.append(event.get("source_document_name"))
        if event.get("source_page"):
            citation_parts.append(f"p. {event.get('source_page')}")
        citation = " | ".join(part for part in citation_parts if part) or event.get("letter_no") or event.get("title")
        row = {
            "source_key": f"S{idx}",
            "source_id": str(event.get("_id")),
            "source_type": "chronology_event",
            "allowed_use": event.get("pleading_use") if event.get("pleading_use") != "none" else "chronology",
            "label": event.get("title") or "Chronology event",
            "citation": citation,
            "snippet": _shorten(event.get("description") or event.get("manual_notes"), 800),
            "page_numbers": [event["source_page"]] if event.get("source_page") else [],
            "clause_number": ", ".join(event.get("contract_clauses") or []) or None,
            "letter_no": event.get("letter_no"),
            "metadata": {
                "chronology_id": event.get("chronology_id"),
                "event_classification": event.get("event_classification"),
                "supports_party": event.get("supports_party"),
                "verification_status": event.get("verification_status"),
                "confidence_score": event.get("confidence_score"),
                "issue_tags": event.get("issue_tags") or [],
                "claim_heads": event.get("claim_heads") or [],
            },
            "source_hash": "",
        }
        row["source_hash"] = _source_hash(row)
        return row

    async def _append_revision(
        self,
        event: Dict[str, Any],
        action: ChronologyRevisionAction,
        current_user: Any,
        *,
        before: Optional[Dict[str, Any]] = None,
        after: Optional[Dict[str, Any]] = None,
        note: Optional[str] = None,
    ) -> Dict[str, Any]:
        latest = await self.db.matter_chronology_event_revisions.find_one(
            {"chronology_id": event["chronology_id"], "event_id": event["_id"]},
            sort=[("revision", -1)],
        )
        revision = MatterChronologyEventRevision(
            chronology_id=event["chronology_id"],
            event_id=str(event["_id"]),
            revision=int((latest or {}).get("revision") or 0) + 1,
            action=action,
            before=before,
            after=after,
            note=note,
            created_by=_actor_id(current_user),
        ).model_dump(by_alias=True)
        await self.db.matter_chronology_event_revisions.insert_one(_jsonable(revision))
        return revision

    async def _refresh_counts(self, chronology_id: str) -> None:
        rows = await _collect(self.db.matter_chronology_events.find({"chronology_id": chronology_id}))
        counts: Dict[str, int] = {"total": len(rows)}
        for row in rows:
            key = str(row.get("verification_status") or "unknown")
            counts[key] = counts.get(key, 0) + 1
        await self.db.matter_chronologies.update_one(
            {"_id": chronology_id},
            {"$set": {"summary_counts": counts, "updated_at": datetime.utcnow()}},
        )

    async def _emit(
        self,
        action: str,
        chronology: Dict[str, Any],
        current_user: Any,
        *,
        before: Optional[Dict[str, Any]] = None,
        after: Optional[Dict[str, Any]] = None,
    ) -> None:
        await self.audit.emit(
            action=f"chronology.{action}",
            actor_id=_actor_id(current_user),
            resource_type="matter_chronology",
            resource_id=str(chronology.get("_id")),
            organization_id=chronology.get("organization_id"),
            project_id=chronology.get("project_id"),
            before=before,
            after=after,
        )

    def _csv_bytes(self, events: List[Dict[str, Any]]) -> bytes:
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(["Date", "Title", "Classification", "Status", "Supports", "Source", "Clauses", "Description"])
        for event in events:
            writer.writerow(
                [
                    event.get("date_text") or event.get("event_date") or "",
                    event.get("title") or "",
                    event.get("event_classification") or "",
                    event.get("verification_status") or "",
                    event.get("supports_party") or "",
                    event.get("source_document_name") or event.get("source_document_id") or "",
                    ", ".join(event.get("contract_clauses") or []),
                    event.get("description") or "",
                ]
            )
        return buffer.getvalue().encode("utf-8-sig")

    def _docx_bytes(self, chronology: Dict[str, Any], events: List[Dict[str, Any]]) -> bytes:
        import docx

        document = docx.Document()
        document.add_heading(chronology.get("title") or "Chronology", level=1)
        table = document.add_table(rows=1, cols=6)
        headers = ["Date", "Event", "Classification", "Status", "Supports", "Source"]
        for idx, label in enumerate(headers):
            table.rows[0].cells[idx].text = label
        for event in events:
            row = table.add_row().cells
            row[0].text = str(event.get("date_text") or event.get("event_date") or "")
            row[1].text = str(event.get("title") or "")
            row[2].text = str(event.get("event_classification") or "")
            row[3].text = str(event.get("verification_status") or "")
            row[4].text = str(event.get("supports_party") or "")
            row[5].text = str(event.get("source_document_name") or event.get("source_document_id") or "")
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    def _pdf_bytes(self, chronology: Dict[str, Any], events: List[Dict[str, Any]]) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, title="Chronology")
        styles = getSampleStyleSheet()
        flow = [Paragraph(escape(chronology.get("title") or "Chronology"), styles["Heading1"])]
        for event in events:
            date = event.get("date_text") or str(event.get("event_date") or "")
            flow.append(Paragraph(escape(f"{date} - {event.get('title') or 'Event'}"), styles["Heading3"]))
            flow.append(Paragraph(escape(_shorten(event.get("description"), 1000) or "[Evidence required]"), styles["BodyText"]))
            flow.append(Spacer(1, 6))
        doc.build(flow)
        return buffer.getvalue()
