from __future__ import annotations

import io
from typing import Any, Dict


class ArbitrationDraftExporter:
    @staticmethod
    def build_pdf(version: Dict[str, Any]) -> bytes:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from xml.sax.saxutils import escape

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(buffer, pagesize=A4, title="Arbitration Pleading")
        styles = getSampleStyleSheet()
        flow = []
        for raw in (version.get("full_markdown") or "").splitlines():
            line = raw.rstrip()
            if not line:
                flow.append(Spacer(1, 6))
            elif line.startswith("## "):
                flow.append(Paragraph(escape(line[3:]), styles["Heading2"]))
            elif line.startswith("# "):
                flow.append(Paragraph(escape(line[2:]), styles["Heading1"]))
            else:
                flow.append(Paragraph(escape(line), styles["BodyText"]))
        doc.build(flow)
        return buffer.getvalue()

    @staticmethod
    def build_docx(version: Dict[str, Any]) -> bytes:
        try:
            import docx
        except ModuleNotFoundError:
            return ArbitrationDraftExporter._build_minimal_docx(version.get("full_markdown") or "")

        document = docx.Document()
        for raw in (version.get("full_markdown") or "").splitlines():
            line = raw.rstrip()
            if not line:
                document.add_paragraph("")
            elif line.startswith("## "):
                document.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                document.add_heading(line[2:], level=1)
            else:
                document.add_paragraph(line)
        buffer = io.BytesIO()
        document.save(buffer)
        return buffer.getvalue()

    @staticmethod
    def _build_minimal_docx(markdown: str) -> bytes:
        import zipfile
        from xml.sax.saxutils import escape

        def paragraph(line: str) -> str:
            text = escape(line.strip())
            if line.startswith("# "):
                return f'<w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>{escape(line[2:].strip())}</w:t></w:r></w:p>'
            if line.startswith("## "):
                return f'<w:p><w:pPr><w:pStyle w:val="Heading2"/></w:pPr><w:r><w:t>{escape(line[3:].strip())}</w:t></w:r></w:p>'
            return f'<w:p><w:r><w:t xml:space="preserve">{text}</w:t></w:r></w:p>'

        body = "".join(paragraph(line.rstrip()) for line in markdown.splitlines())
        document_xml = (
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
            '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            f"<w:body>{body}<w:sectPr><w:pgSz w:w=\"11906\" w:h=\"16838\"/><w:pgMar w:top=\"1440\" w:right=\"1440\" w:bottom=\"1440\" w:left=\"1440\"/></w:sectPr></w:body>"
            "</w:document>"
        )
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(
                "[Content_Types].xml",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
                    '<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>'
                    '<Default Extension="xml" ContentType="application/xml"/>'
                    '<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>'
                    "</Types>"
                ),
            )
            archive.writestr(
                "_rels/.rels",
                (
                    '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                    '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                    '<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>'
                    "</Relationships>"
                ),
            )
            archive.writestr("word/document.xml", document_xml)
        return buffer.getvalue()
