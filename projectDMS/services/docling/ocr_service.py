from __future__ import annotations

import asyncio
import re
import shutil
import uuid
from datetime import datetime
from enum import Enum
from pathlib import Path
from typing import Any, Dict, List, Optional

import orjson
import pytesseract
from fastapi import BackgroundTasks, FastAPI, File, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from loguru import logger
from markdownify import markdownify as md_convert
from pydantic import BaseModel, BaseSettings, Field
from qdrant_client import QdrantClient
from qdrant_client.http import models as qmodels
try:
    from sentence_transformers import SentenceTransformer
except Exception:  # pragma: no cover - optional dependency
    SentenceTransformer = None  # type: ignore

from pdfminer.high_level import extract_text
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
from PIL import Image


class Settings(BaseSettings):
    input_dir: Path = Field(Path("/data/input"), env="DOCLING_INPUT_DIR")
    output_dir: Path = Field(Path("/data/output"), env="DOCLING_OUTPUT_DIR")
    log_dir: Path = Field(Path("/app/logs"), env="DOCLING_LOG_DIR")
    allowed_extensions: List[str] = Field(
        default_factory=lambda: [ext.strip().lower() for ext in (".pdf,.png,.jpg,.jpeg,.tiff,.bmp,.docx,.pptx".split(","))],
        env="DOCLING_ALLOWED_EXTENSIONS",
    )
    max_file_size_mb: int = Field(150, env="DOCLING_MAX_FILE_SIZE_MB")
    log_level: str = Field("INFO", env="DOCLING_LOG_LEVEL")
    vectordb_url: Optional[str] = Field(None, env="VECTORDB_URL")
    vectordb_api_key: Optional[str] = Field(None, env="VECTORDB_API_KEY")
    vectordb_collection: str = Field("documents", env="VECTORDB_COLLECTION")
    vectordb_batch_size: int = Field(64, env="VECTORDB_BATCH_SIZE")
    embedding_model: str = Field("sentence-transformers/all-MiniLM-L6-v2", env="EMBEDDING_MODEL")

    class Config:
        env_file = ".env"
        env_file_encoding = "utf-8"

    def model_post_init(self, __context: Any) -> None:
        if isinstance(self.allowed_extensions, str):
            self.allowed_extensions = [ext.strip().lower() for ext in self.allowed_extensions.split(",") if ext]


settings = Settings()
settings.log_dir.mkdir(parents=True, exist_ok=True)
settings.input_dir.mkdir(parents=True, exist_ok=True)
settings.output_dir.mkdir(parents=True, exist_ok=True)

logger.remove()
logger.add(settings.log_dir / "docling.log", rotation="10 MB", retention="15 files", level=settings.log_level)
logger.add(lambda msg: print(msg, end=""), level=settings.log_level)

app = FastAPI(title="Docling Document Processor", version="0.2.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


class JobStatus(str, Enum):
    PENDING = "pending"
    PROCESSING = "processing"
    COMPLETED = "completed"
    FAILED = "failed"


class DocumentStatus(str, Enum):
    PENDING = "pending"
    PROCESSING = "processing"
    COMPLETED = "completed"
    FAILED = "failed"


class OutputFormat(str, Enum):
    TEXT = "text"
    JSON = "json"
    MARKDOWN = "markdown"


class DocumentResult(BaseModel):
    filename: str
    status: DocumentStatus
    outputs: Dict[str, str] = Field(default_factory=dict)
    metadata: Dict[str, Any] = Field(default_factory=dict)
    error: Optional[str] = None


class JobResult(BaseModel):
    job_id: str
    status: JobStatus
    created_at: datetime
    updated_at: datetime
    total_files: int
    processed_files: int
    results: List[DocumentResult]


class VectorClient:
    def __init__(self) -> None:
        self.client: Optional[QdrantClient] = None
        self.embedder: Optional[SentenceTransformer] = None
        self.dim: Optional[int] = None

        if settings.vectordb_url and SentenceTransformer is not None:
            try:
                self.client = QdrantClient(url=settings.vectordb_url, api_key=settings.vectordb_api_key)
                self.embedder = SentenceTransformer(settings.embedding_model)
                self.dim = self.embedder.get_sentence_embedding_dimension()
                self._ensure_collection()
                logger.info("Vector client initialised for collection '%s'", settings.vectordb_collection)
            except Exception as exc:  # pragma: no cover - optional path
                logger.error("Failed to initialise vector client: %s", exc)
                self.client = None
                self.embedder = None
                self.dim = None
        elif settings.vectordb_url:
            logger.warning("SentenceTransformer not available; vector integration disabled")

    def _ensure_collection(self) -> None:
        if not self.client or not self.dim:
            return
        try:
            collections = self.client.get_collections().collections
            if not any(col.name == settings.vectordb_collection for col in collections):
                logger.info("Creating Qdrant collection '%s'", settings.vectordb_collection)
                self.client.create_collection(
                    collection_name=settings.vectordb_collection,
                    vectors_config=qmodels.VectorParams(size=self.dim, distance=qmodels.Distance.COSINE),
                )
        except Exception as exc:
            logger.error("Unable to ensure Qdrant collection: %s", exc)

    def upsert(self, doc_id: str, text: str, metadata: Dict[str, Any]) -> None:
        if not self.client or not self.embedder:
            return
        try:
            vector = self.embedder.encode(text, convert_to_numpy=False).tolist()
            self.client.upsert(
                collection_name=settings.vectordb_collection,
                points=[qmodels.PointStruct(id=doc_id, vector=vector, payload=metadata)],
            )
        except Exception as exc:  # pragma: no cover - external dependency
            logger.error("Failed to upsert vector for %s: %s", doc_id, exc)


vector_client = VectorClient()

jobs: Dict[str, Dict[str, Any]] = {}
jobs_lock = asyncio.Lock()

INVALID_CHARS = re.compile(r"[^A-Za-z0-9\-_.]")


def sanitize_filename(filename: str) -> str:
    cleaned = INVALID_CHARS.sub("_", filename)
    return cleaned or f"file_{uuid.uuid4().hex}"


def clean_text(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def extract_pdf(path: Path) -> Dict[str, Any]:
    try:
        text = extract_text(str(path)) or ""
    except Exception as exc:
        logger.debug("PDF text extraction failed via pdfminer: %s", exc)
        text = ""
    try:
        reader = PdfReader(str(path))
        metadata = {k: str(v) for k, v in (reader.metadata or {}).items()}
        structure = [clean_text(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:
        logger.debug("PDF parsing failed via PyPDF: %s", exc)
        metadata = {}
        structure = []
    return {
        "text": clean_text(text),
        "metadata": metadata,
        "structure": structure,
    }


def extract_docx(path: Path) -> Dict[str, Any]:
    doc = Document(str(path))
    paragraphs = [clean_text(p.text) for p in doc.paragraphs if p.text]
    props = doc.core_properties
    metadata = {
        "author": props.author,
        "category": props.category,
        "comments": props.comments,
        "content_status": props.content_status,
        "created": props.created.isoformat() if props.created else None,
        "identifier": props.identifier,
        "keywords": props.keywords,
        "language": props.language,
        "last_modified_by": props.last_modified_by,
        "last_printed": props.last_printed.isoformat() if props.last_printed else None,
        "modified": props.modified.isoformat() if props.modified else None,
        "revision": props.revision,
        "subject": props.subject,
        "title": props.title,
        "version": getattr(props, "version", None),
    }
    metadata = {k: v for k, v in metadata.items() if v}
    return {
        "text": clean_text("\n".join(paragraphs)),
        "metadata": metadata,
        "structure": paragraphs,
    }


def extract_pptx(path: Path) -> Dict[str, Any]:
    prs = Presentation(str(path))
    slides: List[List[str]] = []
    for slide_idx, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text:
                texts.append(clean_text(shape.text))
        slides.append(texts)
    flat_text = clean_text("\n".join("\n".join(slide) for slide in slides))
    metadata = {
        "slide_count": len(prs.slides),
        "author": getattr(prs.core_properties, "author", None),
        "created": prs.core_properties.created.isoformat() if prs.core_properties.created else None,
    }
    metadata = {k: v for k, v in metadata.items() if v}
    return {
        "text": flat_text,
        "metadata": metadata,
        "structure": slides,
    }


def extract_image(path: Path) -> Dict[str, Any]:
    with Image.open(path) as img:
        text = pytesseract.image_to_string(img)
        metadata = {"width": img.width, "height": img.height, "mode": img.mode}
    return {
        "text": clean_text(text),
        "metadata": metadata,
        "structure": [],
    }


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".pptx": extract_pptx,
    ".png": extract_image,
    ".jpg": extract_image,
    ".jpeg": extract_image,
    ".tiff": extract_image,
    ".bmp": extract_image,
}


def write_outputs(job_id: str, base_name: str, result: Dict[str, Any], formats: List[OutputFormat]) -> Dict[str, str]:
    output_paths: Dict[str, str] = {}
    job_dir = settings.output_dir / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    text_content = result["text"]
    json_payload = {
        "text": result["text"],
        "metadata": result["metadata"],
        "structure": result["structure"],
    }
    markdown_content = md_convert(result["text"]) if result["text"] else ""

    if OutputFormat.TEXT in formats:
        path = job_dir / f"{base_name}.txt"
        path.write_text(text_content, encoding="utf-8")
        output_paths[OutputFormat.TEXT.value] = str(path)
    if OutputFormat.JSON in formats:
        path = job_dir / f"{base_name}.json"
        path.write_bytes(orjson.dumps(json_payload, option=orjson.OPT_INDENT_2))
        output_paths[OutputFormat.JSON.value] = str(path)
    if OutputFormat.MARKDOWN in formats:
        path = job_dir / f"{base_name}.md"
        path.write_text(markdown_content, encoding="utf-8")
        output_paths[OutputFormat.MARKDOWN.value] = str(path)

    return output_paths


async def process_documents(job_id: str, files: List[Path], sanitized_names: List[str], original_names: List[str], formats: List[OutputFormat]) -> None:
    async with jobs_lock:
        job = jobs[job_id]
        job["status"] = JobStatus.PROCESSING
        job["updated_at"] = datetime.utcnow()

    for idx, (src_path, safe_name, original_name) in enumerate(zip(files, sanitized_names, original_names)):
        ext = src_path.suffix.lower()
        item = job["results"][idx]
        base_name = Path(safe_name).stem
        try:
            item["status"] = DocumentStatus.PROCESSING
            if ext not in EXTRACTORS:
                raise ValueError(f"Unsupported file extension: {ext}")
            extractor = EXTRACTORS[ext]
            data = extractor(src_path)
            outputs = write_outputs(job_id, base_name, data, formats)
            item.update(
                {
                    "status": DocumentStatus.COMPLETED,
                    "outputs": outputs,
                    "metadata": data["metadata"],
                    "error": None,
                }
            )
            if data["text"]:
                vector_client.upsert(f"{job_id}:{safe_name}", data["text"], {"filename": original_name})
        except Exception as exc:
            logger.exception("Failed processing %s", original_name)
            item.update(
                {
                    "status": DocumentStatus.FAILED,
                    "error": str(exc),
                }
            )
        finally:
            async with jobs_lock:
                job = jobs[job_id]
                job["processed_files"] += 1
                job["updated_at"] = datetime.utcnow()

    async with jobs_lock:
        job = jobs[job_id]
        if any(res["status"] == DocumentStatus.FAILED for res in job["results"]):
            job["status"] = JobStatus.FAILED
        else:
            job["status"] = JobStatus.COMPLETED
        job["updated_at"] = datetime.utcnow()


@app.get("/health")
def health() -> Dict[str, str]:
    return {"status": "ok"}


@app.post("/jobs", response_model=JobResult)
async def submit_job(
    background: BackgroundTasks,
    files: List[UploadFile] = File(...),
    output_formats: Optional[str] = None,
) -> JobResult:
    if not files:
        raise HTTPException(status_code=400, detail="At least one file is required")

    formats = [OutputFormat.JSON, OutputFormat.TEXT]
    if output_formats:
        try:
            formats = [OutputFormat(fmt.strip()) for fmt in output_formats.split(",") if fmt.strip()]
        except ValueError as exc:
            raise HTTPException(status_code=400, detail=f"Invalid output format: {exc}") from exc

    job_id = uuid.uuid4().hex
    now = datetime.utcnow()
    dest_files: List[Path] = []
    sanitized_names: List[str] = []
    original_names: List[str] = []

    for upload in files:
        suffix = Path(upload.filename or "").suffix.lower()
        if suffix not in settings.allowed_extensions:
            raise HTTPException(status_code=400, detail=f"Unsupported file extension: {suffix}")
        upload.file.seek(0, 2)
        size_mb = upload.file.tell() / (1024 * 1024)
        upload.file.seek(0)
        if size_mb > settings.max_file_size_mb:
            raise HTTPException(status_code=400, detail=f"File {upload.filename} exceeds max size {settings.max_file_size_mb} MB")

        original_name = upload.filename or uuid.uuid4().hex
        safe_name = sanitize_filename(original_name)
        dest_path = settings.input_dir / job_id / safe_name
        dest_path.parent.mkdir(parents=True, exist_ok=True)
        with dest_path.open("wb") as buffer:
            shutil.copyfileobj(upload.file, buffer)
        dest_files.append(dest_path)
        sanitized_names.append(safe_name)
        original_names.append(original_name)

    async with jobs_lock:
        jobs[job_id] = {
            "job_id": job_id,
            "status": JobStatus.PENDING,
            "created_at": now,
            "updated_at": now,
            "total_files": len(dest_files),
            "processed_files": 0,
            "results": [
                {
                    "filename": name,
                    "status": DocumentStatus.PENDING,
                    "outputs": {},
                    "metadata": {},
                    "error": None,
                }
                for name in original_names
            ],
        }

    background.add_task(process_documents, job_id, dest_files, sanitized_names, original_names, formats)
    return await get_job(job_id)


@app.get("/jobs/{job_id}", response_model=JobResult)
async def get_job(job_id: str) -> JobResult:
    async with jobs_lock:
        job = jobs.get(job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found")
        return JobResult(
            job_id=job["job_id"],
            status=job["status"],
            created_at=job["created_at"],
            updated_at=job["updated_at"],
            total_files=job["total_files"],
            processed_files=job["processed_files"],
            results=[DocumentResult(**res) for res in job["results"]],
        )

